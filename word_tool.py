"""
Word form tool — fills blank table cells in the JM Commercial Application form.

Strategy:
  1. Convert .doc → .docx via LibreOffice if needed
  2. Walk every table cell; match label in previous cell, fill next blank cell
  3. Also supports {{PLACEHOLDER}} replacement for future templates
"""

import os
import re
import subprocess
from docx import Document

BLANK_PATTERN = re.compile(r"^\s*$")


def _convert_doc_to_docx(doc_path: str) -> str:
    out_dir = os.path.dirname(os.path.abspath(doc_path))
    result = subprocess.run(
        ["soffice", "--headless", "--convert-to", "docx", doc_path, "--outdir", out_dir],
        capture_output=True, text=True
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")
    base = os.path.splitext(os.path.basename(doc_path))[0]
    return os.path.join(out_dir, base + ".docx")


def _cell_text(cell) -> str:
    return "".join(p.text for p in cell.paragraphs).strip()


def _set_cell_text(cell, value: str):
    for para in cell.paragraphs:
        if para.runs:
            para.runs[0].text = value
            for run in para.runs[1:]:
                run.text = ""
            return
        else:
            para.add_run(value)
            return


def _fill_by_label_map(doc: Document, label_map: dict) -> dict:
    results = {k: "not_found" for k in label_map}
    norm = {k.lower().strip().rstrip(":"): (k, v) for k, v in label_map.items()}

    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            for i, cell in enumerate(cells):
                text = _cell_text(cell).lower().strip().rstrip(":").strip()
                matched_key, matched_val = None, None
                if text in norm:
                    matched_key, matched_val = norm[text]
                else:
                    for nk, (ok, ov) in norm.items():
                        if nk and text and (nk in text or text in nk):
                            matched_key, matched_val = ok, ov
                            break

                if matched_key and matched_val is not None:
                    for j in range(i + 1, len(cells)):
                        ct = _cell_text(cells[j])
                        if BLANK_PATTERN.match(ct):
                            _set_cell_text(cells[j], str(matched_val))
                            results[matched_key] = "filled"
                            break

    return results


def _fill_placeholders(doc: Document, fields: dict) -> list:
    replaced = []

    def replace_in_para(para):
        full = "".join(r.text for r in para.runs)
        new = full
        for k, v in fields.items():
            ph = f"{{{{{k}}}}}"
            if ph in new:
                new = new.replace(ph, str(v))
                replaced.append(k)
        if new != full and para.runs:
            para.runs[0].text = new
            for r in para.runs[1:]:
                r.text = ""

    for para in doc.paragraphs:
        replace_in_para(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_para(para)
    return replaced


def modify_word_form(template_path: str, output_path: str, fields: dict) -> dict:
    if not os.path.isfile(template_path):
        return {"error": f"Template not found: {template_path}"}

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    work_path = template_path
    if template_path.lower().endswith(".doc"):
        try:
            work_path = _convert_doc_to_docx(template_path)
        except Exception as e:
            return {"error": f"Conversion failed: {e}"}

    try:
        doc = Document(work_path)
    except Exception as e:
        return {"error": f"Failed to open document: {e}"}

    label_results = _fill_by_label_map(doc, fields)
    placeholder_results = _fill_placeholders(doc, fields)

    try:
        doc.save(output_path)
    except Exception as e:
        return {"error": f"Failed to save: {e}"}

    filled = [k for k, v in label_results.items() if v == "filled"] + placeholder_results
    not_found = [k for k, v in label_results.items() if v == "not_found" and k not in placeholder_results]

    return {
        "status": "success",
        "output_path": output_path,
        "fields_filled": list(set(filled)),
        "fields_not_found": list(set(not_found)),
        "total_fields": len(fields),
    }
